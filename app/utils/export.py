import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Font, Alignment, PatternFill


def export_to_excel(queryset, headers, get_row, filename):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_alignment = Alignment(horizontal="center")

    worksheet.append(headers)

    for col in range(1, len(headers) + 1):
        cell = worksheet.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    alt_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    for row_index, obj in enumerate(queryset, start=2):
        worksheet.append(get_row(obj))

        if row_index % 2 == 0:
            for col_index in range(1, len(headers) + 1):
                worksheet.cell(row=row_index, column=col_index).fill = alt_fill

    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = max_length + 2

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = f'attachment; filename={filename}'

    workbook.save(response)
    return response
